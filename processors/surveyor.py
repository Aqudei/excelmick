# https://sbq.com.au/find-a-surveyor/search-sbq-registrant-list/?type=all&search-type=name&title=Paul+Leonard+Heald&postcode=&radius=10&q=Search

from bs4 import BeautifulSoup
from playwright.async_api import Page
import requests
import logging
from db import upsert, find_by_name
import openpyxl
from datetime import datetime

logger = logging.getLogger(__name__)

DB_NAME = "surveyor"


async def fetch_item(name: str):
    url = "https://sbq.com.au/find-a-surveyor/search-sbq-registrant-list/"
    # "?type=all&search-type=name&title=Paul+Leonard+Heald&postcode=&radius=10&q=Search"
    params = {
        "type": "all",
        "search-type": "name",
        "title": name,
        "postcode": "",
        "radius": 10,
        "q": "Search",
    }
    headers = {
        "referer": "https://sbq.com.au/find-a-surveyor/search-sbq-registrant-list/",
        "user-agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/144.0.0.0 Safari/537.36",
    }
    response = requests.get(url, params=params, headers=headers, timeout=60)

    soup = BeautifulSoup(response.text, "html.parser")
    if soup.select_one("div.search-message"):
        return

    for item in soup.select("div.search_result"):
        # Name and role
        h4 = item.find("h4")
        name = h4.contents[0].strip() if h4 else None
        role = h4.find("span").get_text(strip=True) if h4 else None

        # Details table
        details = {}
        for row in item.select(".details table tr"):
            key = row.find("strong").get_text(strip=True)
            value = row.find_all("td")[1].get_text(" ", strip=True)
            details[key] = value

        # Types (Cadastral, Consulting, etc.)
        types = [t.get_text(strip=True) for t in item.select(".types span")]

        result = {
            "name": name,
            "role": role,
            "reg_no": details.get("Rego No."),
            "phone": details.get("Phone"),
            "email": details.get("Email"),
            "address": details.get("Address"),
            "types": types,
        }
        return result


async def boaq_search_reg_no(page: Page, reg_no):
    pass


async def handle_sheet(df, sheet_name, input_file):
    logger.info("Fetching surveyor updates...")
    await fetch_updates(df, sheet_name)

    logger.info("Applying surveyor updates...")
    await apply_updates(sheet_name, input_file)


async def apply_updates(sheet_name, input_file):
    wb = openpyxl.load_workbook(input_file)
    ws = wb[sheet_name]

    value_column = 6
    timestamp_column = 7
    reg_no_column = 8

    for row_idx, _row in enumerate(ws.iter_rows(min_row=2), start=1):
        fullname = (
            f"{(_row[1].value or '').strip()} {(_row[0].value or '').strip()}".strip()
        )
        if fullname in [None, ""]:
            continue

        surveyor = find_by_name(DB_NAME, fullname)

        ws.cell(row=row_idx + 1, column=value_column).value = surveyor.get("status")
        ws.cell(row=row_idx + 1, column=timestamp_column).value = datetime.now()
        ws.cell(row=row_idx + 1, column=reg_no_column).value = surveyor.get(
            "reg_no", ""
        )

    wb.save(filename=input_file)
    wb.close()


async def fetch_updates(df, sheet_name):
    logger.info("Processing 'surveyor' sheet <{}>...".format(sheet_name))

    for index, row in df.iterrows():
        fullname = f"{(row.first_name or '').strip('\r\n\t ')} {(row.surname or '').strip('\r\n\t ')}"
        info = await fetch_item(fullname)
        if not info:
            upsert(DB_NAME, {"status": "Not Found", "name": fullname}, fullname)

        else:
            upsert(DB_NAME, {**info, "status": "Active"}, fullname)
